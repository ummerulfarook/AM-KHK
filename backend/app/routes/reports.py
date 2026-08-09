"""
Reports blueprint — full implementation of Profit & Loss analytics and Excel exporter (Milestone 6).
"""
import io
from datetime import datetime, date, timedelta, timezone
from flask import Blueprint, request, jsonify, send_file
from flask_login import login_required
from sqlalchemy import func
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from app import db
from app.models.sale import RetailSale, SaleItem
from app.models.order import WholesaleOrder, WholesaleOrderItem, PurchaseOrder
from app.models.expense import Expense
from app.models.product import Product
from app.models.customer import Customer
from app.models.credit import CreditLedger, Payment
from app.utils.role_guard import require_roles

reports_bp = Blueprint("reports", __name__, url_prefix="/api/reports")


def local_date_to_utc_range(target_date: date) -> tuple[datetime, datetime]:
    """Convert a local date (IST) to the starting and ending UTC datetimes."""
    # Start of local day is 18:30 of the previous day in UTC
    local_start = datetime.combine(target_date, datetime.min.time())
    utc_start = (local_start - timedelta(hours=5, minutes=30)).replace(tzinfo=timezone.utc)
    
    # End of local day is 18:29:59 of the target day in UTC
    local_end = datetime.combine(target_date, datetime.max.time())
    utc_end = (local_end - timedelta(hours=5, minutes=30)).replace(tzinfo=timezone.utc)
    
    return utc_start, utc_end


def _get_date_filters():
    date_from = request.args.get("dateFrom")
    date_to = request.args.get("dateTo")

    start_date = None
    end_date = None

    if date_from:
        try:
            # Parse YYYY-MM-DD
            parsed_from = date.fromisoformat(date_from)
            start_date, _ = local_date_to_utc_range(parsed_from)
        except ValueError:
            try:
                # Fallback if datetime string is sent
                dt = datetime.fromisoformat(date_from)
                start_date, _ = local_date_to_utc_range(dt.date())
            except ValueError:
                pass
    if date_to:
        try:
            parsed_to = date.fromisoformat(date_to)
            _, end_date = local_date_to_utc_range(parsed_to)
        except ValueError:
            try:
                dt = datetime.fromisoformat(date_to)
                _, end_date = local_date_to_utc_range(dt.date())
            except ValueError:
                pass

    return start_date, end_date


@reports_bp.route("/dashboard", methods=["GET"])
@login_required
@require_roles("owner", "manager", "accountant")
def get_reports_dashboard():
    """Retrieve financial report aggregates: POS vs Wholesale, COGS, Net Profit, category expenses, and sales trends."""
    start_date, end_date = _get_date_filters()

    # 1. Retail Sales Revenue & Items for COGS
    rs_stmt = db.select(RetailSale)
    if start_date:
        rs_stmt = rs_stmt.where(RetailSale.created_at >= start_date)
    if end_date:
        rs_stmt = rs_stmt.where(RetailSale.created_at <= end_date)

    retail_sales = db.session.execute(rs_stmt).scalars().all()
    retail_revenue = sum(s.total for s in retail_sales)
    
    # Calculate POS COGS
    retail_cogs = 0
    for sale in retail_sales:
        for item in sale.items:
            prod = item.product
            # COGS = quantity * product purchase price
            retail_cogs += int(item.quantity * (prod.purchase_price if prod else 0))

    # 2. Wholesale Sales Revenue & Items for COGS (Exclude cancelled)
    wo_stmt = db.select(WholesaleOrder).where(WholesaleOrder.status != "cancelled")
    if start_date:
        wo_stmt = wo_stmt.where(WholesaleOrder.created_at >= start_date)
    if end_date:
        wo_stmt = wo_stmt.where(WholesaleOrder.created_at <= end_date)

    wholesale_orders = db.session.execute(wo_stmt).scalars().all()
    wholesale_revenue = sum(o.total_amount for o in wholesale_orders)

    # Calculate Wholesale COGS
    wholesale_cogs = 0
    for order in wholesale_orders:
        for item in order.items:
            prod = item.product
            wholesale_cogs += int(item.quantity * (prod.purchase_price if prod else 0))

    # Total financial calculations
    total_revenue = retail_revenue + wholesale_revenue
    total_cogs = retail_cogs + wholesale_cogs
    gross_profit = total_revenue - total_cogs
    gross_margin = (gross_profit / total_revenue * 100) if total_revenue > 0 else 0

    # Cash/UPI/Bank sales (mode-wise split for Retail POS, total-mode for Wholesale)
    cash_sales = sum(s.cash_received for s in retail_sales) + sum(o.total_amount for o in wholesale_orders if o.payment_method == "cash")
    upi_sales = sum(s.upi_received for s in retail_sales) + sum(o.total_amount for o in wholesale_orders if o.payment_method == "upi")
    bank_sales = sum(s.bank_received for s in retail_sales) + sum(o.total_amount for o in wholesale_orders if o.payment_method == "bank")
    
    # credit sales represents the portion of sales not paid during checkout
    credit_sales = sum(s.total - s.amount_paid for s in retail_sales) + sum(o.total_amount for o in wholesale_orders if o.payment_method == "credit")

    # Payments received
    pay_stmt = db.select(Payment)
    if start_date:
        pay_stmt = pay_stmt.where(Payment.recorded_at >= start_date)
    if end_date:
        pay_stmt = pay_stmt.where(Payment.recorded_at <= end_date)
    payments = db.session.execute(pay_stmt).scalars().all()
    
    # Exclude checkout downpayments to prevent double counting
    def is_downpayment(p):
        return p.notes and ("[POS Downpayment]" in p.notes or "[POS Credit Downpayment]" in p.notes)

    cash_payments = sum(p.amount for p in payments if p.method == "cash" and not is_downpayment(p))
    upi_payments = sum(p.amount for p in payments if p.method == "upi" and not is_downpayment(p))
    bank_payments = sum(p.amount for p in payments if p.method == "bank" and not is_downpayment(p))

    cash_received = cash_sales + cash_payments
    upi_received = upi_sales + upi_payments
    bank_received = bank_sales + bank_payments
    total_amount_came = cash_received + upi_received + bank_received

    # 3. Expenses Aggregation (Approved only)
    exp_stmt = db.select(Expense).where(Expense.status == "approved")
    if start_date:
        exp_stmt = exp_stmt.where(Expense.expense_date >= start_date.date())
    if end_date:
        exp_stmt = exp_stmt.where(Expense.expense_date <= end_date.date())

    approved_expenses = db.session.execute(exp_stmt).scalars().all()
    total_expenses = sum(e.amount for e in approved_expenses)

    # 3b. Purchase Orders (Expenses include these purchases)
    po_stmt = db.select(PurchaseOrder).where(PurchaseOrder.status != "cancelled")
    if start_date:
        po_stmt = po_stmt.where(PurchaseOrder.created_at >= start_date)
    if end_date:
        po_stmt = po_stmt.where(PurchaseOrder.created_at <= end_date)
    purchase_orders = db.session.execute(po_stmt).scalars().all()
    purchase_expenses = sum(po.total_amount for po in purchase_orders)

    total_expenses_with_purchases = total_expenses + purchase_expenses
    net_profit = gross_profit - total_expenses_with_purchases

    # 4. Expense Category breakdown
    cat_stmt = db.select(Expense.category, func.sum(Expense.amount)).where(Expense.status == "approved")
    if start_date:
        cat_stmt = cat_stmt.where(Expense.expense_date >= start_date.date())
    if end_date:
        cat_stmt = cat_stmt.where(Expense.expense_date <= end_date.date())
    
    category_summary = db.session.execute(cat_stmt.group_by(Expense.category)).all()
    expense_categories = [{"category": row[0], "value": row[1]} for row in category_summary]

    # Include Purchases category in categories pie if there are any
    if purchase_expenses > 0:
        expense_categories.append({"category": "Inventory Purchases", "value": purchase_expenses})

    # 5. Trend Chart Data (Last 30 Days sales & profit trend)
    trend_data = []
    today = date.today()
    for i in range(29, -1, -1):
        day = today - timedelta(days=i)
        day_start = datetime(day.year, day.month, day.day, 0, 0, 0, tzinfo=timezone.utc)
        day_end = datetime(day.year, day.month, day.day, 23, 59, 59, tzinfo=timezone.utc)

        # POS daily sales
        pos_day = db.session.execute(
            db.select(func.sum(RetailSale.total)).where(RetailSale.created_at.between(day_start, day_end))
        ).scalar() or 0

        # Wholesale daily sales
        wo_day = db.session.execute(
            db.select(func.sum(WholesaleOrder.total_amount))
            .where(WholesaleOrder.status != "cancelled")
            .where(WholesaleOrder.created_at.between(day_start, day_end))
        ).scalar() or 0

        # Daily expense
        exp_day = db.session.execute(
            db.select(func.sum(Expense.amount))
            .where(Expense.status == "approved")
            .where(Expense.expense_date == day)
        ).scalar() or 0

        total_day_sales = pos_day + wo_day
        trend_data.append({
            "date": day.strftime("%d %b"),
            "sales": total_day_sales / 100,
            "expenses": exp_day / 100,
            "profit": (total_day_sales - exp_day) / 100
        })

    # Partner Sales & Customer Aggregations
    am_revenue = sum(s.total for s in retail_sales if s.partner == "am") + sum(o.total_amount for o in wholesale_orders if o.partner == "am")
    khk_revenue = sum(s.total for s in retail_sales if s.partner == "khk") + sum(o.total_amount for o in wholesale_orders if o.partner == "khk")
    neutral_revenue = sum(s.total for s in retail_sales if s.partner not in ("am", "khk")) + sum(o.total_amount for o in wholesale_orders if o.partner not in ("am", "khk"))

    am_cust_count = db.session.execute(db.select(func.count(Customer.id)).where(Customer.partner == "am")).scalar() or 0
    khk_cust_count = db.session.execute(db.select(func.count(Customer.id)).where(Customer.partner == "khk")).scalar() or 0
    neutral_cust_count = db.session.execute(db.select(func.count(Customer.id)).where(Customer.partner.notin_(["am", "khk"]))).scalar() or 0

    return jsonify({
        "data": {
            "posRevenue": retail_revenue,
            "wholesaleRevenue": wholesale_revenue,
            "totalRevenue": total_revenue,
            "cogs": total_cogs,
            "grossProfit": gross_profit,
            "grossMargin": gross_margin,
            "expenses": total_expenses_with_purchases,
            "generalExpenses": total_expenses,
            "purchaseExpenses": purchase_expenses,
            "netProfit": net_profit,
            "expenseCategories": expense_categories,
            "paymentDetails": {
                "totalCame": total_amount_came,
                "cash": cash_received,
                "upi": upi_received,
                "bank": bank_received,
                "credit": credit_sales
            },
            "trend": trend_data,
            "partnerBreakdown": {
                "am": {"revenue": am_revenue, "customers": am_cust_count},
                "khk": {"revenue": khk_revenue, "customers": khk_cust_count},
                "neutral": {"revenue": neutral_revenue, "customers": neutral_cust_count}
            }
        }
    }), 200


@reports_bp.route("/export", methods=["GET"])
@login_required
@require_roles("owner", "manager", "accountant")
def export_excel_report():
    """Generate and stream a multi-sheet Excel financial report using openpyxl."""
    start_date, end_date = _get_date_filters()

    # Fetch data
    sales_stmt = db.select(RetailSale)
    if start_date:
        sales_stmt = sales_stmt.where(RetailSale.created_at >= start_date)
    if end_date:
        sales_stmt = sales_stmt.where(RetailSale.created_at <= end_date)
    sales = db.session.execute(sales_stmt.order_by(RetailSale.created_at.desc())).scalars().all()

    expenses_stmt = db.select(Expense)
    if start_date:
        expenses_stmt = expenses_stmt.where(Expense.expense_date >= start_date.date())
    if end_date:
        expenses_stmt = expenses_stmt.where(Expense.expense_date <= end_date.date())
    expenses = db.session.execute(expenses_stmt.order_by(Expense.expense_date.desc())).scalars().all()

    inventory = db.session.execute(db.select(Product).order_by(Product.name)).scalars().all()

    # Create Workbook
    wb = openpyxl.Workbook()
    
    # Setup styles
    header_fill = PatternFill(start_color="0E3A2A", end_color="0E3A2A", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Segoe UI", size=10, bold=True)
    regular_font = Font(name="Segoe UI", size=10)
    center_align = Alignment(horizontal="center")
    right_align = Alignment(horizontal="right")

    # SHEET 1: Retail Sales
    ws_sales = wb.active
    ws_sales.title = "Retail Sales"
    ws_sales.append(["Invoice Number", "Date", "Customer", "Payment Method", "Subtotal (₹)", "Discount (₹)", "Total (₹)"])
    
    for sale in sales:
        ws_sales.append([
            sale.invoice_number,
            sale.created_at.strftime("%Y-%m-%d %H:%M"),
            sale.customer.name if sale.customer else "Walk-in",
            sale.payment_method.upper(),
            sale.subtotal / 100,
            sale.discount / 100,
            sale.total / 100
        ])

    # SHEET 2: Expenses
    ws_exp = wb.create_sheet(title="Expenses")
    ws_exp.append(["Date", "Category", "Amount (₹)", "Status", "Logged By", "Notes"])
    for exp in expenses:
        ws_exp.append([
            exp.expense_date.strftime("%Y-%m-%d"),
            exp.category.upper(),
            exp.amount / 100,
            exp.status.upper(),
            exp.recorded_by_user.name if exp.recorded_by_user else "System",
            exp.notes or ""
        ])

    # SHEET 3: Inventory
    ws_inv = wb.create_sheet(title="Current Inventory")
    ws_inv.append(["Product ID", "Product Name", "Unit", "Cost Price (₹)", "Selling Price (₹)", "Current Stock", "Stock Value (₹)"])
    for prod in inventory:
        ws_inv.append([
            prod.id,
            prod.name,
            prod.unit,
            prod.purchase_price / 100,
            prod.selling_price / 100,
            prod.current_stock,
            (prod.current_stock * (prod.purchase_price / 100))
        ])

    # Format header row and set column widths dynamically
    for sheet in wb.worksheets:
        # Style Header
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align

        # Style data rows and adjust width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = regular_font
                if isinstance(cell.value, (int, float)):
                    cell.alignment = right_align
                    # format as currency/decimals
                    cell.number_format = "#,##0.00"

        # Auto width
        for col in sheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # Save to memory stream
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)

    filename = f"AM_KHK_Financial_Report_{date.today().strftime('%Y-%m-%d')}.xlsx"
    return send_file(
        out,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename
    )


@reports_bp.route("/customers", methods=["GET"])
@login_required
@require_roles("owner", "manager", "accountant")
def get_customer_purchases_report():
    start_date, end_date = _get_date_filters()
    search = (request.args.get("search") or "").strip()
    partner = (request.args.get("partner") or "all").lower()
    tx_type = (request.args.get("type") or "all").lower()

    transactions = []

    # 1. Retail Sales
    if tx_type in ("all", "retail"):
        rs_stmt = db.select(RetailSale)
        if start_date:
            rs_stmt = rs_stmt.where(RetailSale.created_at >= start_date)
        if end_date:
            rs_stmt = rs_stmt.where(RetailSale.created_at <= end_date)
        if partner != "all":
            rs_stmt = rs_stmt.where(RetailSale.partner == partner)
        
        sales = db.session.execute(rs_stmt).scalars().all()
        for s in sales:
            c_name = s.customer.name if s.customer else (s.billing_customer_name or "Walk-in")
            c_phone = s.customer.phone if s.customer else s.billing_customer_phone
            
            # Apply search filter
            if search:
                name_match = search.lower() in c_name.lower()
                phone_match = c_phone and (search in c_phone)
                if not (name_match or phone_match):
                    continue

            # Items summary
            items_summary = ", ".join(
                f"{item.product.name if item.product else 'Unknown'} ({item.quantity} {item.product.unit if item.product else ''})"
                for item in s.items
            )

            transactions.append({
                "id": f"retail_{s.id}",
                "type": "Retail (POS)",
                "reference": s.invoice_number or f"POS-{s.id}",
                "date": s.created_at.isoformat(),
                "customerName": c_name,
                "customerPhone": c_phone or "N/A",
                "itemsSummary": items_summary,
                "paymentMethod": s.payment_method,
                "total": s.total,
                "partner": s.partner
            })

    # 2. Wholesale Orders
    if tx_type in ("all", "wholesale"):
        wo_stmt = db.select(WholesaleOrder).where(WholesaleOrder.status != "cancelled")
        if start_date:
            wo_stmt = wo_stmt.where(WholesaleOrder.created_at >= start_date)
        if end_date:
            wo_stmt = wo_stmt.where(WholesaleOrder.created_at <= end_date)
        if partner != "all":
            wo_stmt = wo_stmt.where(WholesaleOrder.partner == partner)

        orders = db.session.execute(wo_stmt).scalars().all()
        for o in orders:
            c_name = o.customer.name if o.customer else "N/A"
            c_phone = o.customer.phone if o.customer else None

            # Apply search filter
            if search:
                name_match = search.lower() in c_name.lower()
                phone_match = c_phone and (search in c_phone)
                if not (name_match or phone_match):
                    continue

            # Items summary
            items_summary = ", ".join(
                f"{item.product.name if item.product else 'Unknown'} ({item.quantity} {item.product.unit if item.product else ''})"
                for item in o.items
            )

            transactions.append({
                "id": f"wholesale_{o.id}",
                "type": "Wholesale",
                "reference": f"WO-{o.id:04d}",
                "date": o.created_at.isoformat(),
                "customerName": c_name,
                "customerPhone": c_phone or "N/A",
                "itemsSummary": items_summary,
                "paymentMethod": o.payment_method or "credit",
                "total": o.total_amount,
                "partner": o.partner
            })

    # Sort by date descending
    transactions.sort(key=lambda x: x["date"], reverse=True)

    total_transactions = len(transactions)
    total_sales_value = sum(tx["total"] for tx in transactions)

    return jsonify({
        "data": transactions,
        "summary": {
            "totalCount": total_transactions,
            "totalAmount": total_sales_value
        }
    }), 200


@reports_bp.route("/query", methods=["GET"])
@login_required
@require_roles("owner", "manager", "accountant")
def get_custom_report():
    """Retrieve filtered transactional tables for daily, monthly, yearly, expenses, credits, and purchase reports."""
    report_type = request.args.get("type", "daily").lower()
    target_date_str = request.args.get("date") # YYYY-MM-DD
    partner = request.args.get("partner", "").strip().lower()
    
    start_date = None
    end_date = None
    
    if target_date_str:
        try:
            target_date = date.fromisoformat(target_date_str)
        except ValueError:
            target_date = date.today()
    else:
        target_date = date.today()

    if report_type == "daily":
        start_date, end_date = local_date_to_utc_range(target_date)
    elif report_type == "monthly":
        start_day = date(target_date.year, target_date.month, 1)
        if target_date.month == 12:
            end_day = date(target_date.year, 12, 31)
        else:
            end_day = date(target_date.year, target_date.month + 1, 1) - timedelta(days=1)
        start_date, _ = local_date_to_utc_range(start_day)
        _, end_date = local_date_to_utc_range(end_day)
    elif report_type == "yearly":
        if target_date.month >= 4:
            fy_start_year = target_date.year
        else:
            fy_start_year = target_date.year - 1
        start_day = date(fy_start_year, 4, 1)
        end_day = date(fy_start_year + 1, 3, 31)
        start_date, _ = local_date_to_utc_range(start_day)
        _, end_date = local_date_to_utc_range(end_day)
    else:
        df = request.args.get("dateFrom")
        dt = request.args.get("dateTo")
        if df:
            try:
                parsed_df = date.fromisoformat(df)
                start_date, _ = local_date_to_utc_range(parsed_df)
            except ValueError:
                try:
                    parsed_df = datetime.fromisoformat(df).date()
                    start_date, _ = local_date_to_utc_range(parsed_df)
                except ValueError:
                    pass
        if dt:
            try:
                parsed_dt = date.fromisoformat(dt)
                _, end_date = local_date_to_utc_range(parsed_dt)
            except ValueError:
                try:
                    parsed_dt = datetime.fromisoformat(dt).date()
                    _, end_date = local_date_to_utc_range(parsed_dt)
                except ValueError:
                    pass

    if report_type in ("daily", "monthly", "yearly"):
        s_stmt = db.select(RetailSale)
        if start_date:
            s_stmt = s_stmt.where(RetailSale.created_at >= start_date)
        if end_date:
            s_stmt = s_stmt.where(RetailSale.created_at <= end_date)
        if partner in ("am", "khk", "neutral"):
            s_stmt = s_stmt.where(RetailSale.partner == partner)
        sales = db.session.execute(s_stmt).scalars().all()
        
        expenses = []
        if not partner or partner == "all":
            e_stmt = db.select(Expense).where(Expense.status == "approved")
            if start_date:
                e_stmt = e_stmt.where(Expense.expense_date >= start_date.date())
            if end_date:
                e_stmt = e_stmt.where(Expense.expense_date <= end_date.date())
            expenses = db.session.execute(e_stmt).scalars().all()
        
        purchases = []
        if not partner or partner == "all":
            p_stmt = db.select(PurchaseOrder).where(PurchaseOrder.status != "cancelled")
            if start_date:
                p_stmt = p_stmt.where(PurchaseOrder.created_at >= start_date)
            if end_date:
                p_stmt = p_stmt.where(PurchaseOrder.created_at <= end_date)
            purchases = db.session.execute(p_stmt).scalars().all()

        # Fetch dues payments recorded in this period and filter by customer partner
        pay_stmt = db.select(Payment).join(CreditLedger).join(Customer)
        if start_date:
            pay_stmt = pay_stmt.where(Payment.recorded_at >= start_date)
        if end_date:
            pay_stmt = pay_stmt.where(Payment.recorded_at <= end_date)
        if partner in ("am", "khk", "neutral"):
            pay_stmt = pay_stmt.where(Customer.partner == partner)
        payments = db.session.execute(pay_stmt).scalars().all()

        rows = []
        for s in sales:
            rows.append({
                "type": "sale",
                "reference": s.invoice_number or f"SALE-{s.id}",
                "date": s.created_at.isoformat(),
                "particulars": f"Retail Sale ({s.payment_method.upper()})",
                "income": s.total,
                "expense": 0,
                "boxes": sum(item.boxes or 0 for item in s.items)
            })
        for e in expenses:
            rows.append({
                "type": "expense",
                "reference": f"EXP-{e.id}",
                "date": datetime(e.expense_date.year, e.expense_date.month, e.expense_date.day).isoformat(),
                "particulars": f"Expense: {e.category} ({e.notes or ''})",
                "income": 0,
                "expense": e.amount,
                "boxes": 0
            })
        for p in purchases:
            rows.append({
                "type": "purchase",
                "reference": f"PO-{p.id}",
                "date": p.created_at.isoformat(),
                "particulars": f"PO Purchase: {p.supplier.name if p.supplier else 'N/A'}",
                "income": 0,
                "expense": p.total_amount,
                "boxes": 0
            })
        for pay in payments:
            # Exclude checkout downpayments to prevent double counting
            if pay.notes and ("[POS Downpayment]" in pay.notes or "[POS Credit Downpayment]" in pay.notes):
                continue
            c_name = pay.credit_entry.customer.name if (pay.credit_entry and pay.credit_entry.customer) else "N/A"
            ref_num = pay.credit_entry.invoice_ref if pay.credit_entry else ""
            particulars = f"Dues Payment: {c_name}"
            if ref_num:
                particulars += f" (For {ref_num})"
            particulars += f" via {pay.method.upper()}"
            
            rows.append({
                "type": "payment",
                "reference": f"PAY-{pay.id}",
                "date": pay.recorded_at.isoformat(),
                "particulars": particulars,
                "income": pay.amount,
                "expense": 0,
                "boxes": 0
            })
            
        rows.sort(key=lambda x: x["date"], reverse=True)
        total_income = sum(r["income"] for r in rows)
        total_expense = sum(r["expense"] for r in rows)
        
        return jsonify({
            "data": rows,
            "summary": {
                "income": total_income,
                "expense": total_expense,
                "net": total_income - total_expense
            }
        }), 200

    elif report_type == "expense":
        e_stmt = db.select(Expense).where(Expense.status == "approved")
        if start_date:
            e_stmt = e_stmt.where(Expense.expense_date >= start_date.date())
        if end_date:
            e_stmt = e_stmt.where(Expense.expense_date <= end_date.date())
        expenses = db.session.execute(e_stmt).scalars().all()

        p_stmt = db.select(PurchaseOrder).where(PurchaseOrder.status != "cancelled")
        if start_date:
            p_stmt = p_stmt.where(PurchaseOrder.created_at >= start_date)
        if end_date:
            p_stmt = p_stmt.where(PurchaseOrder.created_at <= end_date)
        purchases = db.session.execute(p_stmt).scalars().all()

        rows = []
        for e in expenses:
            rows.append({
                "type": "general",
                "date": datetime(e.expense_date.year, e.expense_date.month, e.expense_date.day).isoformat(),
                "category": e.category,
                "reference": f"EXP-{e.id}",
                "notes": e.notes or "",
                "amount": e.amount
            })
        for p in purchases:
            rows.append({
                "type": "purchase_po",
                "date": p.created_at.isoformat(),
                "category": "Inventory Purchase",
                "reference": f"PO-{p.id}",
                "notes": f"Supplier: {p.supplier.name if p.supplier else 'N/A'}",
                "amount": p.total_amount
            })

        rows.sort(key=lambda x: x["date"], reverse=True)
        return jsonify({
            "data": rows,
            "summary": {
                "total": sum(r["amount"] for r in rows)
            }
        }), 200

    elif report_type == "credit":
        stmt = db.select(Customer).where(Customer.outstanding_balance > 0)
        search = request.args.get("search", "").strip()
        if search:
            stmt = stmt.where(Customer.name.ilike(f"%{search}%"))
        if partner in ("am", "khk", "neutral"):
            stmt = stmt.where(Customer.partner == partner)
        customers = db.session.execute(stmt).scalars().all()
        
        rows = [{
            "id": c.id,
            "name": c.name,
            "phone": c.phone or "N/A",
            "type": c.type,
            "outstandingBalance": c.outstanding_balance,
            "partner": c.partner
        } for c in customers]
        
        rows.sort(key=lambda x: x["outstandingBalance"], reverse=True)
        return jsonify({
            "data": rows,
            "summary": {
                "totalOutstanding": sum(c["outstandingBalance"] for c in rows),
                "count": len(rows)
            }
        }), 200

    elif report_type == "purchase":
        stmt = db.select(PurchaseOrder)
        if start_date:
            stmt = stmt.where(PurchaseOrder.created_at >= start_date)
        if end_date:
            stmt = stmt.where(PurchaseOrder.created_at <= end_date)
        purchases = db.session.execute(stmt.order_by(PurchaseOrder.created_at.desc())).scalars().all()
        
        rows = [{
            "id": p.id,
            "supplierName": p.supplier.name if p.supplier else "N/A",
            "date": p.created_at.isoformat(),
            "status": p.status,
            "paymentStatus": p.payment_status,
            "totalAmount": p.total_amount,
            "notes": p.notes or ""
        } for p in purchases]
        
        return jsonify({
            "data": rows,
            "summary": {
                "total": sum(p["totalAmount"] for p in rows),
                "count": len(rows)
            }
        }), 200

    elif report_type == "purchase_payment":
        stmt = db.select(PurchaseOrder).where(PurchaseOrder.payment_status != "pending")
        if start_date:
            stmt = stmt.where(PurchaseOrder.created_at >= start_date)
        if end_date:
            stmt = stmt.where(PurchaseOrder.created_at <= end_date)
        purchases = db.session.execute(stmt.order_by(PurchaseOrder.created_at.desc())).scalars().all()
        
        rows = []
        for p in purchases:
            rows.append({
                "id": p.id,
                "supplierName": p.supplier.name if p.supplier else "N/A",
                "date": p.created_at.isoformat(),
                "paymentStatus": p.payment_status,
                "amount": p.total_amount,
                "notes": p.notes or ""
            })
            
        return jsonify({
            "data": rows,
            "summary": {
                "total": sum(p["amount"] for p in rows),
                "count": len(rows)
            }
        }), 200

    return jsonify({"error": "Invalid report type"}), 400


@reports_bp.route("/upi-summary", methods=["GET"])
@login_required
@require_roles("owner", "manager", "accountant")
def get_upi_summary():
    """Retrieve total UPI payments grouped by account/UPI ID configured in settings."""
    start_date, end_date = _get_date_filters()

    # 1. POS Direct UPI Receipts
    rs_query = db.select(RetailSale).where(RetailSale.payment_method == "upi")
    if start_date:
        rs_query = rs_query.where(RetailSale.created_at >= start_date)
    if end_date:
        rs_query = rs_query.where(RetailSale.created_at <= end_date)
    rs_receipts = db.session.execute(rs_query).scalars().all()

    # 2. Wholesale Direct UPI Receipts
    wo_query = db.select(WholesaleOrder).where(WholesaleOrder.payment_method == "upi").where(WholesaleOrder.status == "delivered")
    if start_date:
        wo_query = wo_query.where(WholesaleOrder.created_at >= start_date)
    if end_date:
        wo_query = wo_query.where(WholesaleOrder.created_at <= end_date)
    wo_receipts = db.session.execute(wo_query).scalars().all()

    # 3. Credit Ledger UPI Payments
    pay_query = db.select(Payment).where(Payment.method == "upi")
    if start_date:
        pay_query = pay_query.where(Payment.recorded_at >= start_date)
    if end_date:
        pay_query = pay_query.where(Payment.recorded_at <= end_date)
    payments_receipts = db.session.execute(pay_query).scalars().all()

    # Group by upi_id
    accounts = {}

    def add_receipt(upi_id, amount, tx_type, ref, date_val):
        key = (upi_id or "Unassigned/Other").strip() or "Unassigned/Other"
        if key not in accounts:
            accounts[key] = {
                "upiId": key,
                "totalAmount": 0,
                "transactions": []
            }
        accounts[key]["totalAmount"] += amount
        accounts[key]["transactions"].append({
            "type": tx_type,
            "ref": ref,
            "amount": amount,
            "date": date_val.isoformat()
        })

    for r in rs_receipts:
        add_receipt(r.upi_id, r.total, "retail", r.invoice_number or f"RS-{r.id}", r.created_at)

    for w in wo_receipts:
        add_receipt(w.upi_id, w.total_amount, "wholesale", f"WO-{w.id:04d}", w.created_at)

    for p in payments_receipts:
        ref = p.credit_entry.invoice_ref if p.credit_entry else f"PM-{p.id}"
        add_receipt(p.upi_id, p.amount, "ledger_payment", ref, p.recorded_at)

    # Sort transactions by date desc for each account
    results = []
    for key, data in accounts.items():
        data["transactions"].sort(key=lambda x: x["date"], reverse=True)
        results.append(data)

    # Sort accounts by totalAmount desc
    results.sort(key=lambda x: x["totalAmount"], reverse=True)

    return jsonify({"data": results}), 200


@reports_bp.route("/bank-summary", methods=["GET"])
@login_required
@require_roles("owner", "manager", "accountant")
def get_bank_summary():
    """Retrieve total Bank Transfer payments grouped by Bank name configured in settings."""
    start_date, end_date = _get_date_filters()

    # 1. POS Direct Bank Receipts
    rs_query = db.select(RetailSale).where(RetailSale.payment_method == "bank")
    if start_date:
        rs_query = rs_query.where(RetailSale.created_at >= start_date)
    if end_date:
        rs_query = rs_query.where(RetailSale.created_at <= end_date)
    rs_receipts = db.session.execute(rs_query).scalars().all()

    # 2. Wholesale Direct Bank Receipts
    wo_query = db.select(WholesaleOrder).where(WholesaleOrder.payment_method == "bank").where(WholesaleOrder.status == "delivered")
    if start_date:
        wo_query = wo_query.where(WholesaleOrder.created_at >= start_date)
    if end_date:
        wo_query = wo_query.where(WholesaleOrder.created_at <= end_date)
    wo_receipts = db.session.execute(wo_query).scalars().all()

    # 3. Credit Ledger Bank Payments
    pay_query = db.select(Payment).where(Payment.method == "bank")
    if start_date:
        pay_query = pay_query.where(Payment.recorded_at >= start_date)
    if end_date:
        pay_query = pay_query.where(Payment.recorded_at <= end_date)
    payments_receipts = db.session.execute(pay_query).scalars().all()

    # Group by bank_name
    accounts = {}

    def add_receipt(bank_name, amount, tx_type, ref, date_val):
        key = (bank_name or "Unassigned/Other").strip() or "Unassigned/Other"
        if key not in accounts:
            accounts[key] = {
                "bankName": key,
                "totalAmount": 0,
                "transactions": []
            }
        accounts[key]["totalAmount"] += amount
        accounts[key]["transactions"].append({
            "type": tx_type,
            "ref": ref,
            "amount": amount,
            "date": date_val.isoformat()
        })

    for r in rs_receipts:
        add_receipt(r.bank_name, r.total, "retail", r.invoice_number or f"RS-{r.id}", r.created_at)

    for w in wo_receipts:
        add_receipt(w.bank_name, w.total_amount, "wholesale", f"WO-{w.id:04d}", w.created_at)

    for p in payments_receipts:
        ref = p.credit_entry.invoice_ref if p.credit_entry else f"PM-{p.id}"
        add_receipt(p.bank_name, p.amount, "ledger_payment", ref, p.recorded_at)

    # Sort transactions by date desc for each account
    results = []
    for key, data in accounts.items():
        data["transactions"].sort(key=lambda x: x["date"], reverse=True)
        results.append(data)

    # Sort accounts by totalAmount desc
    results.sort(key=lambda x: x["totalAmount"], reverse=True)

    return jsonify({"data": results}), 200


@reports_bp.route("/pdf", methods=["GET"])
@login_required
@require_roles("owner", "manager", "accountant")
def download_pdf_report():
    """Generate and stream a PDF version of the custom report."""
    report_type = request.args.get("type", "daily").lower()
    target_date_str = request.args.get("date") # YYYY-MM-DD
    partner = request.args.get("partner", "").strip().lower()
    
    start_date = None
    end_date = None
    
    if target_date_str:
        try:
            target_date = date.fromisoformat(target_date_str)
        except ValueError:
            target_date = date.today()
    else:
        target_date = date.today()

    if report_type == "daily":
        start_date, end_date = local_date_to_utc_range(target_date)
        period_str = target_date.strftime("%d %b %Y")
    elif report_type == "monthly":
        start_day = date(target_date.year, target_date.month, 1)
        if target_date.month == 12:
            end_day = date(target_date.year, 12, 31)
        else:
            end_day = date(target_date.year, target_date.month + 1, 1) - timedelta(days=1)
        start_date, _ = local_date_to_utc_range(start_day)
        _, end_date = local_date_to_utc_range(end_day)
        period_str = target_date.strftime("%B %Y")
    elif report_type == "yearly":
        if target_date.month >= 4:
            fy_start_year = target_date.year
        else:
            fy_start_year = target_date.year - 1
        start_day = date(fy_start_year, 4, 1)
        end_day = date(fy_start_year + 1, 3, 31)
        start_date, _ = local_date_to_utc_range(start_day)
        _, end_date = local_date_to_utc_range(end_day)
        period_str = f"FY {fy_start_year}-{fy_start_year+1}"
    else:
        df = request.args.get("dateFrom")
        dt = request.args.get("dateTo")
        if df:
            try:
                parsed_df = date.fromisoformat(df)
                start_date, _ = local_date_to_utc_range(parsed_df)
            except ValueError:
                try:
                    parsed_df = datetime.fromisoformat(df).date()
                    start_date, _ = local_date_to_utc_range(parsed_df)
                except ValueError:
                    pass
        if dt:
            try:
                parsed_dt = date.fromisoformat(dt)
                _, end_date = local_date_to_utc_range(parsed_dt)
            except ValueError:
                try:
                    parsed_dt = datetime.fromisoformat(dt).date()
                    _, end_date = local_date_to_utc_range(parsed_dt)
                except ValueError:
                    pass
        period_str = f"{df or 'Start'} to {dt or 'End'}"

    # Setup headers and query based on report type
    headers = []
    rows = []
    summary = {}
    title = ""

    def fmt_rupees(paise):
        return '₹' + f"{paise / 100:,.2f}"

    if report_type in ("daily", "monthly", "yearly"):
        p_title = f" ({partner.upper()})" if partner in ("am", "khk") else ""
        title = f"{report_type.capitalize()} Transaction Report{p_title}"
        headers = [
            {"title": "Date", "key": "date_str", "width": "18%"},
            {"title": "Reference", "key": "reference", "width": "15%"},
            {"title": "Particulars", "key": "particulars", "width": "37%"},
            {"title": "Boxes", "key": "boxes", "width": "10%", "align": "right"},
            {"title": "Income", "key": "income_str", "width": "10%", "align": "right"},
            {"title": "Expense", "key": "expense_str", "width": "10%", "align": "right"},
        ]

        s_stmt = db.select(RetailSale)
        if start_date:
            s_stmt = s_stmt.where(RetailSale.created_at >= start_date)
        if end_date:
            s_stmt = s_stmt.where(RetailSale.created_at <= end_date)
        if partner in ("am", "khk", "neutral"):
            s_stmt = s_stmt.where(RetailSale.partner == partner)
        sales = db.session.execute(s_stmt).scalars().all()

        expenses = []
        if not partner or partner == "all":
            e_stmt = db.select(Expense).where(Expense.status == "approved")
            if start_date:
                e_stmt = e_stmt.where(Expense.expense_date >= start_date.date())
            if end_date:
                e_stmt = e_stmt.where(Expense.expense_date <= end_date.date())
            expenses = db.session.execute(e_stmt).scalars().all()

        purchases = []
        if not partner or partner == "all":
            p_stmt = db.select(PurchaseOrder).where(PurchaseOrder.status != "cancelled")
            if start_date:
                p_stmt = p_stmt.where(PurchaseOrder.created_at >= start_date)
            if end_date:
                p_stmt = p_stmt.where(PurchaseOrder.created_at <= end_date)
            purchases = db.session.execute(p_stmt).scalars().all()

        pay_stmt = db.select(Payment).join(CreditLedger).join(Customer)
        if start_date:
            pay_stmt = pay_stmt.where(Payment.recorded_at >= start_date)
        if end_date:
            pay_stmt = pay_stmt.where(Payment.recorded_at <= end_date)
        if partner in ("am", "khk", "neutral"):
            pay_stmt = pay_stmt.where(Customer.partner == partner)
        payments = db.session.execute(pay_stmt).scalars().all()

        for s in sales:
            rows.append({
                "date": s.created_at.isoformat(),
                "date_str": s.created_at.strftime("%d-%m-%Y %H:%M"),
                "reference": s.invoice_number or f"SALE-{s.id}",
                "particulars": f"Retail Sale ({s.payment_method.upper()})",
                "income": s.total,
                "income_str": fmt_rupees(s.total),
                "expense": 0,
                "expense_str": "₹0.00",
                "boxes": sum(item.boxes or 0 for item in s.items)
            })
        for e in expenses:
            d_val = datetime(e.expense_date.year, e.expense_date.month, e.expense_date.day)
            rows.append({
                "date": d_val.isoformat(),
                "date_str": e.expense_date.strftime("%d-%m-%Y"),
                "reference": f"EXP-{e.id}",
                "particulars": f"Expense: {e.category} ({e.notes or ''})",
                "income": 0,
                "income_str": "₹0.00",
                "expense": e.amount,
                "expense_str": fmt_rupees(e.amount),
                "boxes": 0
            })
        for p in purchases:
            rows.append({
                "date": p.created_at.isoformat(),
                "date_str": p.created_at.strftime("%d-%m-%Y %H:%M"),
                "reference": f"PO-{p.id}",
                "particulars": f"PO Purchase: {p.supplier.name if p.supplier else 'N/A'}",
                "income": 0,
                "income_str": "₹0.00",
                "expense": p.total_amount,
                "expense_str": fmt_rupees(p.total_amount),
                "boxes": 0
            })
        for pay in payments:
            if pay.notes and ("[POS Downpayment]" in pay.notes or "[POS Credit Downpayment]" in pay.notes):
                continue
            c_name = pay.credit_entry.customer.name if (pay.credit_entry and pay.credit_entry.customer) else "N/A"
            ref_num = pay.credit_entry.invoice_ref if pay.credit_entry else ""
            particulars = f"Dues Payment: {c_name}"
            if ref_num:
                particulars += f" (For {ref_num})"
            particulars += f" via {pay.method.upper()}"
            
            rows.append({
                "date": pay.recorded_at.isoformat(),
                "date_str": pay.recorded_at.strftime("%d-%m-%Y %H:%M"),
                "reference": f"PAY-{pay.id}",
                "particulars": particulars,
                "income": pay.amount,
                "income_str": fmt_rupees(pay.amount),
                "expense": 0,
                "expense_str": "₹0.00",
                "boxes": 0
            })
            
        rows.sort(key=lambda x: x["date"], reverse=True)
        total_income = sum(r["income"] for r in rows)
        total_expense = sum(r["expense"] for r in rows)
        summary = {
            "Total Income": fmt_rupees(total_income),
            "Total Expenses": fmt_rupees(total_expense),
            "Net Profit": fmt_rupees(total_income - total_expense)
        }

    elif report_type == "expense":
        title = "Expense & Purchase Outflow Report"
        headers = [
            {"title": "Date", "key": "date_str", "width": "18%"},
            {"title": "Category", "key": "category", "width": "22%"},
            {"title": "Reference", "key": "reference", "width": "15%"},
            {"title": "Particulars / Notes", "key": "notes", "width": "35%"},
            {"title": "Amount", "key": "amount_str", "width": "10%", "align": "right"},
        ]

        expenses = db.session.execute(
            db.select(Expense)
            .where(Expense.status == "approved")
            .where(Expense.expense_date >= start_date.date() if start_date else True)
            .where(Expense.expense_date <= end_date.date() if end_date else True)
        ).scalars().all()

        purchases = db.session.execute(
            db.select(PurchaseOrder)
            .where(PurchaseOrder.status != "cancelled")
            .where(PurchaseOrder.created_at >= start_date if start_date else True)
            .where(PurchaseOrder.created_at <= end_date if end_date else True)
        ).scalars().all()

        for e in expenses:
            d_val = datetime(e.expense_date.year, e.expense_date.month, e.expense_date.day)
            rows.append({
                "date": d_val.isoformat(),
                "date_str": e.expense_date.strftime("%d-%m-%Y"),
                "category": e.category,
                "reference": f"EXP-{e.id}",
                "notes": e.notes or "",
                "amount": e.amount,
                "amount_str": fmt_rupees(e.amount)
            })
        for p in purchases:
            rows.append({
                "date": p.created_at.isoformat(),
                "date_str": p.created_at.strftime("%d-%m-%Y %H:%M"),
                "category": "Inventory Purchase",
                "reference": f"PO-{p.id}",
                "notes": f"Supplier: {p.supplier.name if p.supplier else 'N/A'}",
                "amount": p.total_amount,
                "amount_str": fmt_rupees(p.total_amount)
            })

        rows.sort(key=lambda x: x["date"], reverse=True)
        summary = {
            "Total Outflow": fmt_rupees(sum(r["amount"] for r in rows))
        }

    elif report_type == "credit":
        p_title = f" ({partner.upper()})" if partner in ("am", "khk") else ""
        title = f"Customer Outstanding Dues Report{p_title}"
        headers = [
            {"title": "Customer Name", "key": "name", "width": "35%"},
            {"title": "Phone", "key": "phone", "width": "20%"},
            {"title": "Type", "key": "type_str", "width": "15%"},
            {"title": "Partner Block", "key": "partner", "width": "15%"},
            {"title": "Outstanding Dues", "key": "balance_str", "width": "15%", "align": "right", "bold": True},
        ]

        stmt = db.select(Customer).where(Customer.outstanding_balance > 0)
        search = request.args.get("search", "").strip()
        if search:
            stmt = stmt.where(Customer.name.ilike(f"%{search}%"))
        if partner in ("am", "khk", "neutral"):
            stmt = stmt.where(Customer.partner == partner)
        customers = db.session.execute(stmt).scalars().all()
        
        for c in customers:
            rows.append({
                "name": c.name,
                "phone": c.phone or "N/A",
                "type_str": c.type.upper(),
                "partner": c.partner.upper(),
                "balance": c.outstanding_balance,
                "balance_str": fmt_rupees(c.outstanding_balance)
            })
        
        rows.sort(key=lambda x: x["balance"], reverse=True)
        summary = {
            "Total Outstanding Dues": fmt_rupees(sum(r["balance"] for r in rows)),
            "Debtor Accounts Count": f"{len(rows)} parties"
        }

    elif report_type == "purchase":
        title = "Supplier Purchase Orders Report"
        headers = [
            {"title": "Date", "key": "date_str", "width": "18%"},
            {"title": "Supplier", "key": "supplierName", "width": "32%"},
            {"title": "Status", "key": "status", "width": "15%"},
            {"title": "Payment Status", "key": "paymentStatus", "width": "15%"},
            {"title": "Total Amount", "key": "totalAmount_str", "width": "20%", "align": "right"},
        ]

        purchases = db.session.execute(
            db.select(PurchaseOrder)
            .where(PurchaseOrder.created_at >= start_date if start_date else True)
            .where(PurchaseOrder.created_at <= end_date if end_date else True)
            .order_by(PurchaseOrder.created_at.desc())
        ).scalars().all()
        
        for p in purchases:
            rows.append({
                "supplierName": p.supplier.name if p.supplier else "N/A",
                "date_str": p.created_at.strftime("%d-%m-%Y %H:%M"),
                "status": p.status.upper(),
                "paymentStatus": p.payment_status.upper(),
                "amount": p.total_amount,
                "totalAmount_str": fmt_rupees(p.total_amount),
            })
        
        summary = {
            "Total Purchase Value": fmt_rupees(sum(r["amount"] for r in rows)),
            "Order Count": f"{len(rows)} orders"
        }

    elif report_type == "purchase_payment":
        title = "Supplier Purchase Payments Report"
        headers = [
            {"title": "Date", "key": "date_str", "width": "18%"},
            {"title": "Supplier", "key": "supplierName", "width": "35%"},
            {"title": "PO Status", "key": "paymentStatus", "width": "17%"},
            {"title": "Paid Amount", "key": "amount_str", "width": "30%", "align": "right"},
        ]

        purchases = db.session.execute(
            db.select(PurchaseOrder)
            .where(PurchaseOrder.payment_status != "pending")
            .where(PurchaseOrder.created_at >= start_date if start_date else True)
            .where(PurchaseOrder.created_at <= end_date if end_date else True)
            .order_by(PurchaseOrder.created_at.desc())
        ).scalars().all()
        
        for p in purchases:
            rows.append({
                "supplierName": p.supplier.name if p.supplier else "N/A",
                "date_str": p.created_at.strftime("%d-%m-%Y %H:%M"),
                "paymentStatus": p.payment_status.upper(),
                "amount": p.total_amount,
                "amount_str": fmt_rupees(p.total_amount),
            })
            
        summary = {
            "Total Payments Outflow": fmt_rupees(sum(r["amount"] for r in rows)),
            "Payments Count": f"{len(rows)} operations"
        }

    # Render Jinja template to PDF
    from flask import render_template
    import os
    roboto_font_path = os.path.abspath("backend/app/static/fonts/Roboto-Regular.ttf").replace("\\", "/")
    html_content = render_template(
        "report_pdf.html",
        title=title,
        generated_at=datetime.now(timezone(timedelta(hours=5, minutes=30))).strftime("%d-%m-%Y %H:%M"),
        period=period_str,
        customer_filter=None,
        headers=headers,
        rows=rows,
        summary=summary,
        fmtRupees=fmt_rupees,
        roboto_font_path=roboto_font_path
    )

    from app.services.invoice_service import generate_invoice_pdf
    try:
        pdf_bytes = generate_invoice_pdf(html_content)
        return send_file(
            io.BytesIO(pdf_bytes),
            mimetype="application/pdf",
            as_attachment=True,
            download_name=f"{report_type}_report_{datetime.now().strftime('%Y%m%d%H%M')}.pdf"
        )
    except Exception as e:
        return jsonify({"error": f"Failed to generate PDF report: {e}"}), 500


# ── General Product Sales Report Endpoints ─────────────────────────────────────

def _get_product_sales_report_data(date_from, date_to, product_id, store_id):
    # 1. Retail Sales Items
    stmt_rs = db.select(SaleItem).join(RetailSale)
    if date_from:
        stmt_rs = stmt_rs.where(RetailSale.created_at >= f"{date_from} 00:00:00")
    if date_to:
        stmt_rs = stmt_rs.where(RetailSale.created_at <= f"{date_to} 23:59:59")
    if product_id:
        stmt_rs = stmt_rs.where(SaleItem.product_id == int(product_id))
    if store_id:
        stmt_rs = stmt_rs.where(RetailSale.store_id == int(store_id))
    
    sale_items = db.session.execute(stmt_rs).scalars().all()

    # 2. Wholesale Orders Items
    stmt_wo = db.select(WholesaleOrderItem).join(WholesaleOrder)
    if date_from:
        stmt_wo = stmt_wo.where(WholesaleOrder.created_at >= f"{date_from} 00:00:00")
    if date_to:
        stmt_wo = stmt_wo.where(WholesaleOrder.created_at <= f"{date_to} 23:59:59")
    if product_id:
        stmt_wo = stmt_wo.where(WholesaleOrderItem.product_id == int(product_id))
    if store_id:
        stmt_wo = stmt_wo.where(WholesaleOrder.store_id == int(store_id))

    wo_items = db.session.execute(stmt_wo).scalars().all()

    # Aggregate by Product ID
    products_data = {}

    for item in sale_items:
        p_id = item.product_id
        if p_id not in products_data:
            products_data[p_id] = {
                "productName": item.product.name,
                "unit": item.product.unit,
                "quantitySold": 0.0,
                "totalSales": 0,
                "invoiceNumbers": set()
            }
        products_data[p_id]["quantitySold"] += item.quantity
        products_data[p_id]["totalSales"] += item.subtotal
        products_data[p_id]["invoiceNumbers"].add(item.sale.invoice_number or f"SALE-{item.sale.id}")

    for item in wo_items:
        p_id = item.product_id
        if p_id not in products_data:
            products_data[p_id] = {
                "productName": item.product.name,
                "unit": item.product.unit,
                "quantitySold": 0.0,
                "totalSales": 0,
                "invoiceNumbers": set()
            }
        products_data[p_id]["quantitySold"] += item.quantity
        products_data[p_id]["totalSales"] += item.subtotal
        products_data[p_id]["invoiceNumbers"].add(f"WO-{item.order.id:04d}")

    # Convert to list
    report_rows = []
    for p_id, data in products_data.items():
        inv_count = len(data["invoiceNumbers"])
        report_rows.append({
            "productId": p_id,
            "productName": data["productName"],
            "unit": data["unit"],
            "quantitySold": data["quantitySold"],
            "invoicesCount": inv_count,
            "totalSales": data["totalSales"]
        })

    # Sort by productName
    report_rows.sort(key=lambda x: x["productName"])
    return report_rows


@reports_bp.route("/product-sales", methods=["GET"])
@login_required
@require_roles("owner", "manager", "accountant")
def get_general_product_sales_report():
    date_from = request.args.get("dateFrom")
    date_to = request.args.get("dateTo")
    product_id = request.args.get("productId")
    store_id = request.args.get("storeId")

    rows = _get_product_sales_report_data(date_from, date_to, product_id, store_id)

    grand_total_qty = sum(x["quantitySold"] for x in rows)
    grand_total_sales = sum(x["totalSales"] for x in rows)

    return jsonify({
        "data": rows,
        "summary": {
            "grandTotalQuantity": grand_total_qty,
            "grandTotalSales": grand_total_sales
        }
    }), 200


@reports_bp.route("/product-sales-pdf", methods=["GET"])
@login_required
@require_roles("owner", "manager", "accountant")
def get_general_product_sales_report_pdf():
    date_from = request.args.get("dateFrom")
    date_to = request.args.get("dateTo")
    product_id = request.args.get("productId")
    store_id = request.args.get("storeId")

    rows = _get_product_sales_report_data(date_from, date_to, product_id, store_id)

    grand_total_qty = sum(x["quantitySold"] for x in rows)
    grand_total_sales = sum(x["totalSales"] for x in rows)

    period_str = f"{date_from or 'Start'} to {date_to or 'End'}"

    from flask import render_template
    import io
    from app.services.invoice_service import generate_invoice_pdf
    
    def fmt_rupees(paise):
        return '₹' + f"{paise / 100:,.2f}"

    import os
    roboto_font_path = os.path.abspath("backend/app/static/fonts/Roboto-Regular.ttf").replace("\\", "/")
    
    html_content = render_template(
        "product_sales_report_pdf.html",
        generated_at=datetime.now(timezone(timedelta(hours=5, minutes=30))).strftime("%d-%m-%Y %H:%M"),
        period=period_str,
        rows=rows,
        summary={
            "grandTotalQuantity": grand_total_qty,
            "grandTotalSales": grand_total_sales
        },
        fmtRupees=fmt_rupees,
        roboto_font_path=roboto_font_path
    )

    try:
        pdf_bytes = generate_invoice_pdf(html_content)
        return send_file(
            io.BytesIO(pdf_bytes),
            mimetype="application/pdf",
            as_attachment=True,
            download_name=f"product_sales_report_{datetime.now().strftime('%Y%m%d')}.pdf"
        )
    except Exception as e:
        return jsonify({"error": f"Failed to generate product sales report PDF: {e}"}), 500


@reports_bp.route("/product-sales-excel", methods=["GET"])
@login_required
@require_roles("owner", "manager", "accountant")
def get_general_product_sales_report_excel():
    date_from = request.args.get("dateFrom")
    date_to = request.args.get("dateTo")
    product_id = request.args.get("productId")
    store_id = request.args.get("storeId")

    rows = _get_product_sales_report_data(date_from, date_to, product_id, store_id)

    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment
    import io

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Product Sales"

    header_fill = PatternFill(start_color="0E3A2A", end_color="0E3A2A", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Segoe UI", size=10, bold=True)
    regular_font = Font(name="Segoe UI", size=10)

    ws.append(["AM & KHK Vegetable Merchants"])
    ws.append(["General Product Sales Report"])
    ws.append([f"Period: {date_from or 'Start'} to {date_to or 'End'}"])
    ws.append([])

    headers = ["Product Name", "Unit", "Quantity Sold", "Number of Invoices", "Total Sales (₹)"]
    ws.append(headers)

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=5, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font

    for x in rows:
        ws.append([
            x["productName"],
            x["unit"],
            x["quantitySold"],
            x["invoicesCount"],
            x["totalSales"] / 100
        ])

    grand_total_qty = sum(x["quantitySold"] for x in rows)
    grand_total_sales = sum(x["totalSales"] for x in rows)

    ws.append([])
    ws.append(["Grand Totals"])
    ws.cell(row=ws.max_row, column=1).font = bold_font
    
    ws.append(["Grand Total Quantity Sold", grand_total_qty])
    ws.cell(row=ws.max_row, column=1).font = regular_font
    ws.cell(row=ws.max_row, column=2).font = bold_font

    ws.append(["Grand Total Sales Amount (₹)", grand_total_sales / 100])
    ws.cell(row=ws.max_row, column=1).font = regular_font
    ws.cell(row=ws.max_row, column=2).font = bold_font

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f"product_sales_report_{datetime.now().strftime('%Y%m%d')}.xlsx"
    )


