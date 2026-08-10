"""
Customers blueprint — full implementation for CRM (Milestone 4).
"""
from datetime import datetime, date, timedelta, timezone
from decimal import Decimal
from flask import Blueprint, request, jsonify, send_file
from flask_login import login_required, current_user
from sqlalchemy import func, or_
from app import db
from app.models.customer import Customer, CustomerStore
from app.models.sale import RetailSale, SaleItem
from app.models.order import WholesaleOrder, WholesaleOrderItem
from app.models.credit import CreditLedger, Payment
from app.utils.role_guard import require_roles

customers_bp = Blueprint("customers", __name__, url_prefix="/api/customers")


def _rupees_to_paise(val) -> int:
    try:
        return int(Decimal(str(val)) * 100)
    except (TypeError, ValueError, InvalidOperation):
        return 0


# Avoid decimal import error if InvalidOperation is raised
from decimal import InvalidOperation


@customers_bp.route("/", methods=["GET"])
@login_required
def list_customers():
    """List customers with search, type filter, and pagination."""
    page = max(1, request.args.get("page", 1, type=int))
    per_page = min(10000, request.args.get("perPage", 25, type=int))
    search = request.args.get("search", "").strip()
    cust_type = request.args.get("type", "").strip()
    partner = request.args.get("partner", "").strip().lower()

    stmt = db.select(Customer)
    if search:
        stmt = stmt.where(
            or_(
                Customer.name.ilike(f"%{search}%"),
                Customer.phone.ilike(f"%{search}%")
            )
        )
    if cust_type in ("wholesale", "retail"):
        stmt = stmt.where(Customer.type == cust_type)
    if partner in ("am", "khk", "neutral"):
        stmt = stmt.where(Customer.partner == partner)

    total = db.session.execute(
        db.select(func.count()).select_from(stmt.subquery())
    ).scalar() or 0

    customers = db.session.execute(
        stmt.order_by(Customer.name)
        .offset((page - 1) * per_page)
        .limit(per_page)
    ).scalars().all()

    return jsonify({
        "data": [c.to_dict() for c in customers],
        "pagination": {
            "page": page,
            "perPage": per_page,
            "total": total,
            "totalPages": max(1, -(-total // per_page))
        }
    }), 200


@customers_bp.route("/", methods=["POST"])
@login_required
@require_roles("owner", "manager", "accountant")
def create_customer():
    """Create a new customer."""
    data = request.get_json(silent=True) or {}
    name = (data.get("name") or "").strip()
    if not name:
        return jsonify({"error": "Name is required"}), 422

    phone = (data.get("phone") or "").strip() or None
    address = (data.get("address") or "").strip() or None
    cust_type = (data.get("type") or "retail").strip()
    if cust_type not in ("wholesale", "retail"):
        return jsonify({"error": "Type must be 'retail' or 'wholesale'"}), 422

    credit_limit = _rupees_to_paise(data.get("creditLimit", 0))
    opening_balance = _rupees_to_paise(data.get("openingBalance", 0))
    partner = (data.get("partner") or "neutral").strip()

    customer = Customer(
        name=name,
        phone=phone,
        address=address,
        type=cust_type,
        credit_limit=credit_limit,
        opening_balance=opening_balance,
        outstanding_balance=opening_balance,
        partner=partner,
        is_active=1
    )
    db.session.add(customer)
    db.session.flush()

    if opening_balance > 0:
        ledger_entry = CreditLedger(
            customer_id=customer.id,
            invoice_ref="OPENING-BAL",
            amount=opening_balance,
            amount_paid=0,
            due_date=date.today() + timedelta(days=30),
            status="due"
        )
        db.session.add(ledger_entry)

    db.session.commit()

    return jsonify({"data": customer.to_dict(), "message": "Customer created successfully"}), 201


@customers_bp.route("/<int:cust_id>", methods=["GET"])
@login_required
def get_customer(cust_id: int):
    """Get single customer details."""
    customer = db.session.get(Customer, cust_id)
    if not customer:
        return jsonify({"error": "Customer not found"}), 404
    return jsonify({"data": customer.to_dict()}), 200


@customers_bp.route("/<int:cust_id>", methods=["PUT"])
@login_required
@require_roles("owner", "manager", "accountant")
def update_customer(cust_id: int):
    """Update a customer."""
    customer = db.session.get(Customer, cust_id)
    if not customer:
        return jsonify({"error": "Customer not found"}), 404

    data = request.get_json(silent=True) or {}
    name = (data.get("name") or "").strip()
    if not name:
        return jsonify({"error": "Name is required"}), 422

    cust_type = (data.get("type") or "retail").strip()
    if cust_type not in ("wholesale", "retail"):
        return jsonify({"error": "Type must be 'retail' or 'wholesale'"}), 422

    new_ob = _rupees_to_paise(data.get("openingBalance", 0))

    # Sync CreditLedger entry for OPENING-BAL
    ledger_entry = db.session.execute(
        db.select(CreditLedger)
        .where(CreditLedger.customer_id == cust_id)
        .where(CreditLedger.invoice_ref == "OPENING-BAL")
    ).scalar_one_or_none()

    if ledger_entry:
        old_ob = customer.opening_balance
        diff = new_ob - old_ob
        ledger_entry.amount = new_ob
        if ledger_entry.amount_paid >= new_ob:
            ledger_entry.status = "paid"
            ledger_entry.paid_at = datetime.now(timezone.utc)
        else:
            ledger_entry.status = "due"
            ledger_entry.paid_at = None
        customer.outstanding_balance += diff
    else:
        if new_ob > 0:
            ledger_entry = CreditLedger(
                customer_id=customer.id,
                invoice_ref="OPENING-BAL",
                amount=new_ob,
                amount_paid=0,
                due_date=date.today() + timedelta(days=30),
                status="due"
            )
            db.session.add(ledger_entry)
            customer.outstanding_balance += new_ob

    customer.name = name
    customer.phone = (data.get("phone") or "").strip() or None
    customer.address = (data.get("address") or "").strip() or None
    customer.type = cust_type
    customer.credit_limit = _rupees_to_paise(data.get("creditLimit", 0))
    customer.opening_balance = new_ob
    customer.partner = (data.get("partner") or "neutral").strip()

    db.session.commit()
    return jsonify({"data": customer.to_dict(), "message": "Customer updated successfully"}), 200


@customers_bp.route("/<int:cust_id>/toggle-active", methods=["POST"])
@login_required
@require_roles("owner", "manager")
def toggle_customer_active(cust_id: int):
    """Toggle customer active/inactive status."""
    customer = db.session.get(Customer, cust_id)
    if not customer:
        return jsonify({"error": "Customer not found"}), 404

    customer.is_active = 0 if customer.is_active else 1
    db.session.commit()

    status_str = "deactivated" if not customer.is_active else "activated"
    return jsonify({
        "data": customer.to_dict(),
        "message": f"Customer successfully {status_str}"
    }), 200


@customers_bp.route("/<int:cust_id>/history", methods=["GET"])
@login_required
def get_customer_history(cust_id: int):
    """Get customer purchase history (combined RetailSale and WholesaleOrder)."""
    customer = db.session.get(Customer, cust_id)
    if not customer:
        return jsonify({"error": "Customer not found"}), 404

    page = max(1, request.args.get("page", 1, type=int))
    per_page = min(1000, request.args.get("perPage", 10, type=int)) # Support larger page sizes for statement prints
    date_from = request.args.get("dateFrom")
    date_to = request.args.get("dateTo")
    store_id = request.args.get("storeId")

    # Fetch retail sales
    stmt_rs = db.select(RetailSale).where(RetailSale.customer_id == cust_id)
    if date_from:
        stmt_rs = stmt_rs.where(RetailSale.created_at >= f"{date_from} 00:00:00")
    if date_to:
        stmt_rs = stmt_rs.where(RetailSale.created_at <= f"{date_to} 23:59:59")
    if store_id:
        try:
            stmt_rs = stmt_rs.where(RetailSale.store_id == int(store_id))
        except (ValueError, TypeError):
            pass
    retail_sales = db.session.execute(stmt_rs).scalars().all()

    # Fetch wholesale orders
    stmt_wo = db.select(WholesaleOrder).where(WholesaleOrder.customer_id == cust_id)
    if date_from:
        stmt_wo = stmt_wo.where(WholesaleOrder.created_at >= f"{date_from} 00:00:00")
    if date_to:
        stmt_wo = stmt_wo.where(WholesaleOrder.created_at <= f"{date_to} 23:59:59")
    if store_id:
        try:
            stmt_wo = stmt_wo.where(WholesaleOrder.store_id == int(store_id))
        except (ValueError, TypeError):
            pass
    wholesale_orders = db.session.execute(stmt_wo).scalars().all()

    # Combine & format
    history = []
    for rs in retail_sales:
        payment_status = "paid"
        amount_paid = rs.total
        balance = 0
        if rs.payment_method == "credit":
            ledger = db.session.execute(
                db.select(CreditLedger).where(CreditLedger.sale_id == rs.id)
            ).scalars().first()
            if ledger:
                payment_status = ledger.status
                amount_paid = ledger.amount_paid
                balance = ledger.amount - ledger.amount_paid
            else:
                payment_status = "credit"
                amount_paid = 0
                balance = rs.total

        history.append({
            "id": rs.id,
            "type": "retail",
            "invoiceNumber": rs.invoice_number,
            "total": rs.total,
            "itemsCount": len(rs.items),
            "paymentMethod": rs.payment_method,
            "paymentStatus": payment_status,
            "amountPaid": amount_paid,
            "balance": balance,
            "storeId": rs.store_id,
            "storeName": rs.store.name if rs.store else None,
            "createdAt": rs.created_at.isoformat(),
        })

    for wo in wholesale_orders:
        payment_status = "paid"
        amount_paid = wo.total_amount
        balance = 0
        if wo.payment_method == "credit":
            ledger = db.session.execute(
                db.select(CreditLedger).where(CreditLedger.wholesale_order_id == wo.id)
            ).scalars().first()
            if ledger:
                payment_status = ledger.status
                amount_paid = ledger.amount_paid
                balance = ledger.amount - ledger.amount_paid
            else:
                payment_status = "credit"
                amount_paid = 0
                balance = wo.total_amount

        history.append({
            "id": wo.id,
            "type": "wholesale",
            "invoiceNumber": f"WO-{wo.id:04d}",
            "total": wo.total_amount,
            "itemsCount": len(wo.items),
            "paymentMethod": wo.payment_method or "credit",
            "status": wo.status,
            "paymentStatus": payment_status,
            "amountPaid": amount_paid,
            "balance": balance,
            "storeId": wo.store_id,
            "storeName": wo.store.name if wo.store else None,
            "createdAt": wo.created_at.isoformat(),
        })

    # Sort desc by date
    history.sort(key=lambda x: x["createdAt"], reverse=True)

    total = len(history)
    start = (page - 1) * per_page
    end = start + per_page
    paginated_data = history[start:end]

    return jsonify({
        "data": paginated_data,
        "pagination": {
            "page": page,
            "perPage": per_page,
            "total": total,
            "totalPages": max(1, -(-total // per_page))
        }
    }), 200


@customers_bp.route("/<int:cust_id>/credit", methods=["GET"])
@login_required
def get_customer_credit(cust_id: int):
    """Get customer credit summary: limit, outstanding balance, credit entries."""
    customer = db.session.get(Customer, cust_id)
    if not customer:
        return jsonify({"error": "Customer not found"}), 404

    date_from = request.args.get("dateFrom")
    date_to = request.args.get("dateTo")

    stmt_entries = db.select(CreditLedger).where(CreditLedger.customer_id == cust_id)
    if date_from:
        stmt_entries = stmt_entries.where(CreditLedger.created_at >= f"{date_from} 00:00:00")
    if date_to:
        stmt_entries = stmt_entries.where(CreditLedger.created_at <= f"{date_to} 23:59:59")
    entries = db.session.execute(stmt_entries.order_by(CreditLedger.created_at.desc())).scalars().all()

    total_outstanding = customer.outstanding_balance

    stmt_payments = db.select(Payment).join(CreditLedger).where(CreditLedger.customer_id == cust_id)
    if date_from:
        stmt_payments = stmt_payments.where(Payment.recorded_at >= f"{date_from} 00:00:00")
    if date_to:
        stmt_payments = stmt_payments.where(Payment.recorded_at <= f"{date_to} 23:59:59")
    payments = db.session.execute(stmt_payments.order_by(Payment.recorded_at.desc())).scalars().all()

    return jsonify({
        "data": {
            "creditLimit": customer.credit_limit,
            "totalOutstanding": total_outstanding,
            "entries": [e.to_dict() for e in entries],
            "payments": [p.to_dict() for p in payments]
        }
    }), 200


@customers_bp.route("/<int:cust_id>/pay", methods=["POST"])
@login_required
@require_roles("owner", "manager", "accountant")
def clear_customer_dues(cust_id: int):
    """Record a general or invoice-specific customer payment."""
    customer = db.session.get(Customer, cust_id)
    if not customer:
        return jsonify({"error": "Customer not found"}), 404

    data = request.get_json(silent=True) or {}
    amount_rs = data.get("amount")
    if amount_rs is None:
        return jsonify({"error": "Amount is required"}), 422

    amount = _rupees_to_paise(amount_rs)
    if amount <= 0:
        return jsonify({"error": "Amount must be greater than zero"}), 422

    method = (data.get("method") or "cash").strip()
    if method not in ("cash", "upi", "bank"):
        return jsonify({"error": "Method must be cash, upi, or bank"}), 422

    notes = (data.get("notes") or "").strip() or None
    upi_id = data.get("upiId")
    bank_name = data.get("bankName")
    invoice_ref = data.get("invoiceRef")

    try:
        from app.utils.date_helper import get_working_date
        working_dt = get_working_date()

        # 1. Update customer's outstanding balance
        customer.outstanding_balance = max(0, customer.outstanding_balance - amount)

        applied_amount = amount

        # 2. If specific invoiceRef is provided, apply payment to that entry first
        if invoice_ref:
            entry = db.session.execute(
                db.select(CreditLedger)
                .where(CreditLedger.customer_id == cust_id)
                .where(CreditLedger.invoice_ref == invoice_ref)
                .where(CreditLedger.status != "paid")
            ).scalar_one_or_none()
            if entry:
                rem = entry.amount - entry.amount_paid
                applied = min(applied_amount, rem)
                
                p = Payment(
                    credit_ledger_id=entry.id,
                    amount=applied,
                    method=method,
                    upi_id=upi_id,
                    bank_name=bank_name,
                    recorded_by_id=current_user.id,
                    notes=f"[Invoice Payment: {invoice_ref}] {notes or ''}".strip(),
                    recorded_at=working_dt,
                )
                db.session.add(p)
                entry.amount_paid += applied
                if entry.amount_paid >= entry.amount:
                    entry.status = "paid"
                    entry.paid_at = working_dt
                    # Update wholesale order payment status if linked
                    if entry.wholesale_order_id:
                        order = db.session.get(WholesaleOrder, entry.wholesale_order_id)
                        if order:
                            order.payment_status = "paid"
                else:
                    # Update wholesale order payment status if linked
                    if entry.wholesale_order_id:
                        order = db.session.get(WholesaleOrder, entry.wholesale_order_id)
                        if order:
                            order.payment_status = "partial"
                
                applied_amount -= applied

        # 3. Distribute remaining payment across unpaid CreditLedger entries (oldest first)
        if applied_amount > 0:
            unpaid_entries = db.session.execute(
                db.select(CreditLedger)
                .where(CreditLedger.customer_id == cust_id)
                .where(CreditLedger.status != "paid")
                .order_by(CreditLedger.due_date.asc())
            ).scalars().all()

            for entry in unpaid_entries:
                if applied_amount <= 0:
                    break

                rem = entry.amount - entry.amount_paid
                if applied_amount >= rem:
                    # Fully pay this entry
                    p = Payment(
                        credit_ledger_id=entry.id,
                        amount=rem,
                        method=method,
                        upi_id=upi_id,
                        bank_name=bank_name,
                        recorded_by_id=current_user.id,
                        notes=f"[General Payment] {notes or ''}".strip(),
                        recorded_at=working_dt,
                    )
                    db.session.add(p)
                    entry.amount_paid += rem
                    entry.status = "paid"
                    entry.paid_at = working_dt
                    applied_amount -= rem

                    # Update wholesale order payment status if linked
                    if entry.wholesale_order_id:
                        order = db.session.get(WholesaleOrder, entry.wholesale_order_id)
                        if order:
                            order.payment_status = "paid"
                else:
                    # Partially pay this entry
                    p = Payment(
                        credit_ledger_id=entry.id,
                        amount=applied_amount,
                        method=method,
                        upi_id=upi_id,
                        bank_name=bank_name,
                        recorded_by_id=current_user.id,
                        notes=f"[General Payment] {notes or ''}".strip(),
                        recorded_at=working_dt,
                    )
                    db.session.add(p)
                    entry.amount_paid += applied_amount
                    applied_amount = 0

                    # Update wholesale order payment status if linked
                    if entry.wholesale_order_id:
                        order = db.session.get(WholesaleOrder, entry.wholesale_order_id)
                        if order:
                            order.payment_status = "partial"

        db.session.commit()
        return jsonify({
            "data": customer.to_dict(),
            "message": f"Payment of ₹{amount_rs:.2f} successfully recorded and applied to dues."
        }), 200

    except Exception as exc:
        db.session.rollback()
        return jsonify({"error": f"Failed to record customer payment: {exc}"}), 500


@customers_bp.route("/<int:customer_id>", methods=["DELETE"])
@login_required
@require_roles("owner", "manager")
def delete_customer(customer_id: int):
    """Delete a customer and all their associated transaction history (sales, orders, ledgers)."""
    customer = db.session.get(Customer, customer_id)
    if not customer:
        return jsonify({"error": "Customer not found"}), 404
        
    try:
        # 1. Delete all payments linked to credit ledgers of this customer
        ledger_ids = db.session.execute(
            db.select(CreditLedger.id).where(CreditLedger.customer_id == customer_id)
        ).scalars().all()
        if ledger_ids:
            db.session.execute(
                db.delete(Payment).where(Payment.credit_ledger_id.in_(ledger_ids))
            )
            
        # 2. Delete all credit ledger entries of this customer
        db.session.execute(
            db.delete(CreditLedger).where(CreditLedger.customer_id == customer_id)
        )
        
        # 3. Delete wholesale orders (and their items)
        order_ids = db.session.execute(
            db.select(WholesaleOrder.id).where(WholesaleOrder.customer_id == customer_id)
        ).scalars().all()
        for oid in order_ids:
            db.session.execute(
                db.delete(WholesaleOrderItem).where(WholesaleOrderItem.order_id == oid)
            )
        if order_ids:
            db.session.execute(
                db.delete(WholesaleOrder).where(WholesaleOrder.id.in_(order_ids))
            )
            
        # 4. Delete retail sales (and their items)
        sale_ids = db.session.execute(
            db.select(RetailSale.id).where(RetailSale.customer_id == customer_id)
        ).scalars().all()
        for sid in sale_ids:
            db.session.execute(
                db.delete(SaleItem).where(SaleItem.sale_id == sid)
            )
        if sale_ids:
            db.session.execute(
                db.delete(RetailSale).where(RetailSale.id.in_(sale_ids))
            )
            
        # 5. Delete the customer profile
        db.session.delete(customer)
        db.session.commit()
        return jsonify({"message": "Customer and all associated transaction history deleted successfully"}), 200
    except Exception as exc:
        db.session.rollback()
        return jsonify({"error": f"Failed to delete customer: {exc}"}), 500


@customers_bp.route("/<int:cust_id>/ledger-pdf", methods=["GET"])
@login_required
def download_customer_ledger_pdf(cust_id: int):
    """Generate and stream a PDF customer statement ledger."""
    customer = db.session.get(Customer, cust_id)
    if not customer:
        return jsonify({"error": "Customer not found"}), 404

    date_from = request.args.get("dateFrom")
    date_to = request.args.get("dateTo")

    # Fetch retail sales
    stmt_rs = db.select(RetailSale).where(RetailSale.customer_id == cust_id)
    if date_from:
        stmt_rs = stmt_rs.where(RetailSale.created_at >= f"{date_from} 00:00:00")
    if date_to:
        stmt_rs = stmt_rs.where(RetailSale.created_at <= f"{date_to} 23:59:59")
    retail_sales = db.session.execute(stmt_rs).scalars().all()

    # Fetch wholesale orders
    stmt_wo = db.select(WholesaleOrder).where(WholesaleOrder.customer_id == cust_id)
    if date_from:
        stmt_wo = stmt_wo.where(WholesaleOrder.created_at >= f"{date_from} 00:00:00")
    if date_to:
        stmt_wo = stmt_wo.where(WholesaleOrder.created_at <= f"{date_to} 23:59:59")
    wholesale_orders = db.session.execute(stmt_wo).scalars().all()

    # Combine & format
    history = []
    for rs in retail_sales:
        payment_status = "paid"
        amount_paid = rs.total
        balance = 0
        if rs.payment_method == "credit":
            ledger = db.session.execute(
                db.select(CreditLedger).where(CreditLedger.sale_id == rs.id)
            ).scalars().first()
            if ledger:
                payment_status = ledger.status
                amount_paid = ledger.amount_paid
                balance = ledger.amount - ledger.amount_paid
            else:
                payment_status = "credit"
                amount_paid = 0
                balance = rs.total

        history.append({
            "date": rs.created_at.strftime("%d-%m-%Y %H:%M"),
            "type": "retail",
            "invoiceNumber": rs.invoice_number or f"RS-{rs.id}",
            "total": rs.total,
            "paymentMethod": rs.payment_method,
            "paymentStatus": payment_status,
            "amountPaid": amount_paid,
            "balance": balance,
            "createdAt": rs.created_at
        })

    for wo in wholesale_orders:
        payment_status = "paid"
        amount_paid = wo.total_amount
        balance = 0
        if wo.payment_method == "credit":
            ledger = db.session.execute(
                db.select(CreditLedger).where(CreditLedger.wholesale_order_id == wo.id)
            ).scalars().first()
            if ledger:
                payment_status = ledger.status
                amount_paid = ledger.amount_paid
                balance = ledger.amount - ledger.amount_paid
            else:
                payment_status = "credit"
                amount_paid = 0
                balance = wo.total_amount

        history.append({
            "date": wo.created_at.strftime("%d-%m-%Y %H:%M"),
            "type": "wholesale",
            "invoiceNumber": f"WO-{wo.id:04d}",
            "total": wo.total_amount,
            "paymentMethod": wo.payment_method or "credit",
            "paymentStatus": payment_status,
            "amountPaid": amount_paid,
            "balance": balance,
            "createdAt": wo.created_at
        })

    # Sort desc by date
    history.sort(key=lambda x: x["createdAt"], reverse=True)

    period_str = f"{date_from or 'Start'} to {date_to or 'End'}"
    
    # Render PDF statement
    from flask import render_template
    import io
    from app.services.invoice_service import generate_invoice_pdf
    
    def fmt_rupees(paise):
        return '₹' + f"{paise / 100:,.2f}"

    import os
    roboto_font_path = os.path.abspath("backend/app/static/fonts/Roboto-Regular.ttf").replace("\\", "/")
    html_content = render_template(
        "customer_ledger_pdf.html",
        customer=customer,
        generated_at=datetime.now(timezone(timedelta(hours=5, minutes=30))).strftime("%d-%m-%Y %H:%M"),
        period=period_str,
        rows=history,
        fmtRupees=fmt_rupees,
        roboto_font_path=roboto_font_path
    )

    try:
        pdf_bytes = generate_invoice_pdf(html_content)
        return send_file(
            io.BytesIO(pdf_bytes),
            mimetype="application/pdf",
            as_attachment=True,
            download_name=f"statement_{customer.name.replace(' ', '_')}_{datetime.now().strftime('%Y%m%d')}.pdf"
        )
    except Exception as e:
        return jsonify({"error": f"Failed to generate customer ledger PDF: {e}"}), 500


# ── Customer Sub-Stores CRUD routes ──────────────────────────────────────────

@customers_bp.route("/stores", methods=["GET"])
@login_required
def list_all_customer_stores():
    stores = db.session.execute(
        db.select(CustomerStore).where(CustomerStore.is_active == 1).order_by(CustomerStore.name)
    ).scalars().all()
    return jsonify({"data": [{
        "id": s.id,
        "name": s.name,
        "customerName": s.customer.name
    } for s in stores]}), 200


@customers_bp.route("/<int:customer_id>/stores", methods=["GET"])
@login_required
def list_customer_stores(customer_id):
    customer = db.session.get(Customer, customer_id)
    if not customer:
        return jsonify({"error": "Customer not found"}), 404
    stores = db.session.execute(
        db.select(CustomerStore)
        .where(CustomerStore.customer_id == customer_id)
        .order_by(CustomerStore.name)
    ).scalars().all()
    return jsonify({"data": [s.to_dict() for s in stores]}), 200


@customers_bp.route("/<int:customer_id>/stores", methods=["POST"])
@login_required
@require_roles("owner", "manager", "accountant")
def create_customer_store(customer_id):
    customer = db.session.get(Customer, customer_id)
    if not customer:
        return jsonify({"error": "Customer not found"}), 404
    data = request.get_json(silent=True) or {}
    name = (data.get("name") or "").strip()
    if not name:
        return jsonify({"error": "Store name is required"}), 422
    
    # check for duplicate names under same customer
    existing = db.session.execute(
        db.select(CustomerStore).where(
            CustomerStore.customer_id == customer_id,
            CustomerStore.name == name
        )
    ).scalar_one_or_none()
    if existing:
        return jsonify({"error": "Store name already exists for this customer"}), 422

    store = CustomerStore(customer_id=customer_id, name=name, is_active=1)
    db.session.add(store)
    db.session.commit()
    return jsonify({"data": store.to_dict()}), 201


@customers_bp.route("/stores/<int:store_id>", methods=["PUT"])
@login_required
@require_roles("owner", "manager", "accountant")
def update_customer_store(store_id):
    store = db.session.get(CustomerStore, store_id)
    if not store:
        return jsonify({"error": "Store not found"}), 404
    data = request.get_json(silent=True) or {}
    
    if "name" in data:
        name = (data.get("name") or "").strip()
        if not name:
            return jsonify({"error": "Store name cannot be empty"}), 422
        existing = db.session.execute(
            db.select(CustomerStore).where(
                CustomerStore.customer_id == store.customer_id,
                CustomerStore.name == name,
                CustomerStore.id != store_id
            )
        ).scalar_one_or_none()
        if existing:
            return jsonify({"error": "Store name already exists for this customer"}), 422
        store.name = name

    if "isActive" in data:
        store.is_active = 1 if data.get("isActive") else 0

    db.session.commit()
    return jsonify({"data": store.to_dict()}), 200


# ── Customer Sales Report endpoints ──────────────────────────────────────────

def _get_customer_sales_report_data(customer_id, date_from, date_to, product_id, store_id):
    # 1. Retail Sales Items
    stmt_rs = db.select(SaleItem).join(RetailSale).where(RetailSale.customer_id == customer_id)
    if date_from:
        stmt_rs = stmt_rs.where(RetailSale.created_at >= f"{date_from} 00:00:00")
    if date_to:
        stmt_rs = stmt_rs.where(RetailSale.created_at <= f"{date_to} 23:59:59")
    if product_id:
        stmt_rs = stmt_rs.where(SaleItem.product_id == int(product_id))
    if store_id:
        stmt_rs = stmt_rs.where(RetailSale.store_id == int(store_id))
    
    sale_items = db.session.execute(stmt_rs).scalars().all()

    # 2. Wholesale Orders Items
    stmt_wo = db.select(WholesaleOrderItem).join(WholesaleOrder).where(WholesaleOrder.customer_id == customer_id)
    if date_from:
        stmt_wo = stmt_wo.where(WholesaleOrder.created_at >= f"{date_from} 00:00:00")
    if date_to:
        stmt_wo = stmt_wo.where(WholesaleOrder.created_at <= f"{date_to} 23:59:59")
    if product_id:
        stmt_wo = stmt_wo.where(WholesaleOrderItem.product_id == int(product_id))
    if store_id:
        stmt_wo = stmt_wo.where(WholesaleOrder.store_id == int(store_id))

    wo_items = db.session.execute(stmt_wo).scalars().all()

    report_items = []
    for item in sale_items:
        report_items.append({
            "invoiceNumber": item.sale.invoice_number or f"SALE-{item.sale.id}",
            "date": item.sale.created_at,
            "storeName": item.sale.store.name if item.sale.store else "Direct / Main",
            "storeId": item.sale.store_id,
            "productName": item.product.name,
            "quantity": item.quantity,
            "unit": item.product.unit,
            "unitPrice": item.unit_price,
            "totalAmount": item.subtotal
        })

    for item in wo_items:
        report_items.append({
            "invoiceNumber": f"WO-{item.order.id:04d}",
            "date": item.order.created_at,
            "storeName": item.order.store.name if item.order.store else "Direct / Main",
            "storeId": item.order.store_id,
            "productName": item.product.name,
            "quantity": item.quantity,
            "unit": item.product.unit,
            "unitPrice": item.unit_price,
            "totalAmount": item.subtotal
        })

    report_items.sort(key=lambda x: x["date"], reverse=True)
    return report_items


@customers_bp.route("/<int:customer_id>/sales-report", methods=["GET"])
@login_required
@require_roles("owner", "manager", "accountant")
def get_customer_sales_report(customer_id):
    customer = db.session.get(Customer, customer_id)
    if not customer:
        return jsonify({"error": "Customer not found"}), 404

    date_from = request.args.get("dateFrom")
    date_to = request.args.get("dateTo")
    product_id = request.args.get("productId")
    store_id = request.args.get("storeId")

    items = _get_customer_sales_report_data(customer_id, date_from, date_to, product_id, store_id)

    serialized_items = []
    for x in items:
        serialized_items.append({
            **x,
            "date": x["date"].isoformat()
        })

    total_qty = sum(x["quantity"] for x in items)
    total_amount = sum(x["totalAmount"] for x in items)
    unique_invoices = len(set(x["invoiceNumber"] for x in items))

    store_subtotals = {}
    for x in items:
        s_name = x["storeName"]
        if s_name not in store_subtotals:
            store_subtotals[s_name] = {"quantity": 0.0, "amount": 0}
        store_subtotals[s_name]["quantity"] += x["quantity"]
        store_subtotals[s_name]["amount"] += x["totalAmount"]

    return jsonify({
        "data": serialized_items,
        "summary": {
            "totalInvoices": unique_invoices,
            "totalQuantity": total_qty,
            "totalAmount": total_amount
        },
        "storeSubtotals": store_subtotals
    }), 200


@customers_bp.route("/<int:customer_id>/sales-report-pdf", methods=["GET"])
@login_required
@require_roles("owner", "manager", "accountant")
def get_customer_sales_report_pdf(customer_id):
    customer = db.session.get(Customer, customer_id)
    if not customer:
        return jsonify({"error": "Customer not found"}), 404

    date_from = request.args.get("dateFrom")
    date_to = request.args.get("dateTo")
    product_id = request.args.get("productId")
    store_id = request.args.get("storeId")

    items = _get_customer_sales_report_data(customer_id, date_from, date_to, product_id, store_id)

    formatted_rows = []
    for x in items:
        formatted_rows.append({
            **x,
            "date_str": x["date"].strftime("%d-%m-%Y %H:%M")
        })

    period_str = f"{date_from or 'Start'} to {date_to or 'End'}"
    
    total_qty = sum(x["quantity"] for x in items)
    total_amount = sum(x["totalAmount"] for x in items)
    unique_invoices = len(set(x["invoiceNumber"] for x in items))

    store_subtotals = {}
    for x in items:
        s_name = x["storeName"]
        if s_name not in store_subtotals:
            store_subtotals[s_name] = {"quantity": 0.0, "amount": 0}
        store_subtotals[s_name]["quantity"] += x["quantity"]
        store_subtotals[s_name]["amount"] += x["totalAmount"]

    from flask import render_template
    import io
    from app.services.invoice_service import generate_invoice_pdf
    
    def fmt_rupees(paise):
        return '₹' + f"{paise / 100:,.2f}"

    import os
    roboto_font_path = os.path.abspath("backend/app/static/fonts/Roboto-Regular.ttf").replace("\\", "/")
    
    html_content = render_template(
        "customer_sales_report_pdf.html",
        customer=customer,
        generated_at=datetime.now(timezone(timedelta(hours=5, minutes=30))).strftime("%d-%m-%Y %H:%M"),
        period=period_str,
        rows=formatted_rows,
        summary={
            "totalInvoices": unique_invoices,
            "totalQuantity": total_qty,
            "totalAmount": total_amount
        },
        storeSubtotals=store_subtotals,
        fmtRupees=fmt_rupees,
        roboto_font_path=roboto_font_path
    )

    try:
        pdf_bytes = generate_invoice_pdf(html_content)
        return send_file(
            io.BytesIO(pdf_bytes),
            mimetype="application/pdf",
            as_attachment=True,
            download_name=f"sales_report_{customer.name.replace(' ', '_')}_{datetime.now().strftime('%Y%m%d')}.pdf"
        )
    except Exception as e:
        return jsonify({"error": f"Failed to generate customer sales report PDF: {e}"}), 500


@customers_bp.route("/<int:customer_id>/sales-report-excel", methods=["GET"])
@login_required
@require_roles("owner", "manager", "accountant")
def get_customer_sales_report_excel(customer_id):
    customer = db.session.get(Customer, customer_id)
    if not customer:
        return jsonify({"error": "Customer not found"}), 404

    date_from = request.args.get("dateFrom")
    date_to = request.args.get("dateTo")
    product_id = request.args.get("productId")
    store_id = request.args.get("storeId")

    items = _get_customer_sales_report_data(customer_id, date_from, date_to, product_id, store_id)

    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment
    import io

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales Report"

    header_fill = PatternFill(start_color="0E3A2A", end_color="0E3A2A", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Segoe UI", size=10, bold=True)
    regular_font = Font(name="Segoe UI", size=10)

    ws.append(["AM & KHK Vegetable Merchants"])
    ws.append([f"Customer Sales Report: {customer.name}"])
    ws.append([f"Period: {date_from or 'Start'} to {date_to or 'End'}"])
    ws.append([])

    headers = ["Invoice Number", "Date", "Store Name", "Product Name", "Quantity", "Unit", "Unit Price (₹)", "Total Amount (₹)"]
    ws.append(headers)

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=5, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font

    for x in items:
        ws.append([
            x["invoiceNumber"],
            x["date"].strftime("%Y-%m-%d %H:%M"),
            x["storeName"],
            x["productName"],
            x["quantity"],
            x["unit"],
            x["unitPrice"] / 100,
            x["totalAmount"] / 100
        ])

    total_qty = sum(x["quantity"] for x in items)
    total_amount = sum(x["totalAmount"] for x in items)
    unique_invoices = len(set(x["invoiceNumber"] for x in items))

    ws.append([])
    ws.append(["Summary Totals"])
    ws.cell(row=ws.max_row, column=1).font = bold_font
    
    ws.append(["Total Unique Invoices", unique_invoices])
    ws.cell(row=ws.max_row, column=1).font = regular_font
    ws.cell(row=ws.max_row, column=2).font = bold_font

    ws.append(["Total Quantity Sold", total_qty])
    ws.cell(row=ws.max_row, column=1).font = regular_font
    ws.cell(row=ws.max_row, column=2).font = bold_font

    ws.append(["Total Sales Amount (₹)", total_amount / 100])
    ws.cell(row=ws.max_row, column=1).font = regular_font
    ws.cell(row=ws.max_row, column=2).font = bold_font

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f"sales_report_{customer.name.replace(' ', '_')}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    )

